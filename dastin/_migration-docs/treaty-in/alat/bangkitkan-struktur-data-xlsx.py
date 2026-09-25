# -*- coding: utf-8 -*-
r"""
bangkitkan-struktur-data-xlsx.py -- STRUKTUR DATA MODEL BARU dalam bentuk ERD
yang sama dengan Diagram-Skema-Tabel-TreatyIn-dan-EDM-v2.xlsx.

    masukan : 4-erd-dan-tabel-datar/erd-skema-terpadu.json   37 tabel, 41 relasi
    keluaran: 4-erd-dan-tabel-datar/STRUKTUR-DATA-SKEMA-BARU.xlsx

BENTUKNYA DITIRU, ISINYA TIDAK
    Berkas contohnya -- Diagram-Skema-Tabel-TreatyIn-dan-EDM-v2.xlsx -- adalah
    POTRET SISTEM LAMA dan berbanner begitu. Yang ditiru darinya hanya BENTUK:

        kotak tiga baris          baris 1 nama + kardinalitas, baris 2 PK/FK,
                                  baris 3 asal di sistem lama
        indentasi 4 kolom         satu tingkat kedalaman = 4 kolom ke kanan
        lebar kolom 2.3           dan kotak dilebur 46 kolom
        garis kisi MATI           hierarki terbaca dari indentasi, bukan garis
        penanda BERSAMA           entitas yang muncul di lebih dari satu lembar
        lembar Daftar Relasi      dan lembar Catatan & Batas

    Yang TIDAK ditiru: daftar tabelnya. Contohnya memuat 44 tabel SISTEM LAMA
    berawalan T_; berkas ini memuat 37 entitas MODEL BARU. Bannernya karena itu
    berbunyi terbalik -- RANCANGAN, bukan potret.

PEMBAGIAN LEMBAR, dan kenapa berbeda dari contohnya
    Contohnya terbagi Prop / Non Prop / EDM Prop / EDM Non Prop, sebab di sistem
    lama addendum adalah SALINAN pohon. Di model baru addendum ADALAH sebuah
    VERSI_KONTRAK (G1), sehingga tidak ada pohon EDM tersendiri untuk disalin.
    Yang tersisa sebagai pembeda nyata hanya cabang proporsional versus
    non-proporsional -- satu-satunya tempat kedua sisi benar-benar berpisah
    (STRUKTUR-DATA.md sec 1.3).

APA YANG TIDAK DAPAT DIBUKTIKAN BERKAS INI
    Bahwa himpunannya lengkap. Semesta sec 10 masih kurang 340 properti titik
    buta (L-8). Dan dua entitas -- NILAI_SELISIH, BESARAN_DAPAT_DISESUAIKAN --
    BELUM ADA DDL-nya sama sekali; keduanya bertanda merah, bukan didiamkan.

PENOLAKAN dilaporkan satu baris per sebab.

Pakai:  PYTHONIOENCODING=utf-8 python alat/bangkitkan-struktur-data-xlsx.py
"""
import io, os, json, sys
from collections import Counter, OrderedDict, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

AKAR = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
ERD = os.path.join(AKAR, '4-erd-dan-tabel-datar')
SUMBER = os.path.join(ERD, 'erd-skema-terpadu.json')
KELUAR = os.path.join(ERD, 'STRUKTUR-DATA-SKEMA-BARU.xlsx')
TANGGAL = '25 September 2026'

tolak = Counter()
d = json.load(io.open(SUMBER, encoding='utf-8'))
tabel = d['tabel']

# Tabel acuan TIDAK masuk pohon -- ia dirujuk belasan entitas dan akan
# menggandakan kotaknya. Ia punya lembarnya sendiri, dan rujukannya dicetak di
# baris kedua kotak yang merujuknya.
ACUAN = ['MATA_UANG', 'JENIS_POTONGAN', 'JENIS_REASURANSI', 'BAHAYA',
         'KELOMPOK_TREATY', 'KELAS_BISNIS', 'BESARAN_DAPAT_DISESUAIKAN']

# ══ Pohon: anak per induk, hanya lewat kunci asing yang BUKAN ke tabel acuan ═
anak = defaultdict(list)
acuan_ref = defaultdict(list)
for nm, t in tabel.items():
    for f in t['fk']:
        if f['induk'] in ACUAN:
            acuan_ref[nm].append(f)
        else:
            anak[f['induk']].append((nm, f))

# ══ Lembar: DAFTAR BERNAMA, bukan tebakan dari bentuk nama ═════════════════
# Yang dikecualikan tiap cabang disebut satu per satu, beserta sebabnya.
HANYA_NONPROP = ['BAGIAN', 'NILAI_PREMI_BRUTO', 'NILAI_PREMI_BRUTO_MINIMUM',
                 'NILAI_MDP', 'NILAI_MDP_MINIMUM', 'PEMULIHAN_LIMIT']
HANYA_PROP = ['DETAIL_PROPORSIONAL', 'NILAI_CADANGAN_PREMI']
HANYA_ADJ = ['NILAI_SELISIH']

LEMBAR = OrderedDict()
LEMBAR['Proporsional'] = dict(
    akar=['KONTRAK'], buang=HANYA_NONPROP + HANYA_ADJ,
    ket='cabang proporsional — bagian NuRe adalah ATRIBUT `DETAIL_PROPORSIONAL`, bukan baris')
LEMBAR['Non-proporsional'] = dict(
    akar=['KONTRAK'], buang=HANYA_PROP + HANYA_ADJ,
    ket='cabang non-proporsional — `BAGIAN` hanya ada di sini (STRUKTUR-DATA §1.3)')
# Lembar Penyesuaian SENGAJA SEMPIT: ia menggambar SAMBUNGANNYA, bukan mengulang
# ketiga belas anak bersama yang sudah tergambar di kedua lembar cabang.
SAMBUNGAN_ADJ = ['DOKUMEN_ADDENDUM', 'VERSI_KONTRAK', 'NILAI_SELISIH']
LEMBAR['Penyesuaian'] = dict(
    akar=['DOKUMEN_ADDENDUM'], buang=[], sisakan=SAMBUNGAN_ADJ,
    ket='modul Adjustment — sebuah penyesuaian ADALAH sebuah `VERSI_KONTRAK` (G1); '
        'dokumen memayungi banyak versi, lintas kontrak. Lembar ini SENGAJA sempit: '
        'ia menggambar sambungannya, bukan mengulang anak bersama')
LEMBAR['Acuan'] = dict(akar=list(ACUAN), buang=[], datar=True,
                       ket='tabel acuan — wajib tabel, bukan CHECK (ADR-0038, INV-62)')

# ══ Gaya, diukur dari berkas contohnya ═════════════════════════════════════
LEBAR_KOLOM, LEBAR_KOTAK, INDEN = 2.3, 46, 4
AKAR_BIRU = PatternFill('solid', fgColor='1F4E79')
ANAK_BIRU = PatternFill('solid', fgColor='2E75B6')
MERAH = PatternFill('solid', fgColor='C00000')
KUNING = PatternFill('solid', fgColor='FFF2CC')
MUDA = PatternFill('solid', fgColor='EDF3F9')
HIJAU = PatternFill('solid', fgColor='548235')
F_KEPALA = Font(name='Arial', size=9, bold=True, color='FFFFFF')
F_KUNCI = Font(name='Arial', size=8, bold=True, color='000000')
F_ASAL = Font(name='Arial', size=8, color='595959')
F_JUDUL = Font(name='Arial', size=12, bold=True, color='1F4E79')
F_SUB = Font(name='Arial', size=8, color='595959')
F_PITA = Font(name='Arial', size=10, bold=True, color='FFFFFF')
F_TH = Font(name='Arial', size=8, bold=True, color='FFFFFF')
F_TD = Font(name='Arial', size=8)
KIRI = Alignment(horizontal='left', vertical='center', indent=1)
M = Side(style='medium')

PITA1 = 'RANCANGAN MODEL BARU, %s. BUKAN potret sistem lama.' % TANGGAL
PITA2 = ('Dibangkitkan alat/bangkitkan-struktur-data-xlsx.py dari erd-skema-terpadu.json. '
         'Bentuknya meniru Diagram-Skema-Tabel-TreatyIn-dan-EDM-v2.xlsx; isinya TIDAK.')

wb = Workbook()
wb.remove(wb.active)
di_lembar = defaultdict(list)


def kepala(ws, judul, ket, n):
    for c in range(1, 60):
        ws.cell(1, c).fill = HIJAU
        ws.cell(2, c).fill = KUNING
    a = ws.cell(1, 1, PITA1); a.font = F_PITA; a.alignment = KIRI
    b = ws.cell(2, 1, PITA2); b.font = Font(name='Arial', size=8, bold=True); b.alignment = KIRI
    ws.row_dimensions[1].height = 20
    ws.row_dimensions[2].height = 14
    t = ws.cell(3, 2, judul); t.font = F_JUDUL; t.alignment = KIRI
    s = ws.cell(4, 2, '%d entitas · %s' % (n, ket)); s.font = F_SUB; s.alignment = KIRI
    ws.freeze_panes = 'A5'
    ws.sheet_view.showGridLines = False
    for c in range(1, 60):
        ws.column_dimensions[L(c)].width = LEBAR_KOLOM


def kotak(ws, r, c, baris3, warna_kepala):
    """Tiga baris terlebur, bertepi medium -- bentuk yang sama dengan contohnya."""
    for i, (teks, fill, font) in enumerate(baris3):
        rr = r + i
        ws.merge_cells(start_row=rr, start_column=c, end_row=rr, end_column=c + LEBAR_KOTAK - 1)
        sel = ws.cell(rr, c, teks)
        sel.fill = fill
        sel.font = font
        sel.alignment = KIRI
        for cc in range(c, c + LEBAR_KOTAK):
            s = ws.cell(rr, cc)
            s.border = Border(left=M if cc == c else None,
                              right=M if cc == c + LEBAR_KOTAK - 1 else None,
                              top=M if i == 0 else None,
                              bottom=M if i == 2 else None)
    ws.row_dimensions[r].height = 13
    ws.row_dimensions[r + 2].height = 15
    return r + 4


def baris_kotak(nm, fk, akar, lain):
    t = tabel[nm]
    kard = '' if akar else ('1:1' if kardinal(nm, fk) else '1:N')
    tanda = []
    if kard:
        tanda.append(kard)
    tanda.append('%d kolom' % len(t['kolom']))
    if t.get('belum_ddl'):
        tanda.append('BELUM BER-DDL')
    if lain:
        tanda.append('BERSAMA → ' + ', '.join(lain))
    l1 = '%s        %s' % (nm, ' · '.join(tanda))

    if akar or not fk:
        l2 = 'PK   %s' % (t['pk'] or '(tidak terbaca)')
    else:
        od = fk.get('on_delete') or '(tidak dinyatakan)'
        l2 = 'FK   %s → %s.%s     %s' % (fk['kolom'], fk['induk'], fk['kolom_induk'], od)
        l2 += '        PK   %s' % (t['pk'] or '(tidak terbaca)')
    ref = acuan_ref.get(nm) or []
    if ref:
        l2 += '   ·   acuan: ' + ', '.join('%s → %s' % (f['kolom'], f['induk']) for f in ref)

    asal = ' | '.join(t.get('asal_pega') or []) or '(tidak ada di pohon Pega — entitas BARU)'
    lama = ' | '.join(t.get('nama_lama') or [])
    ka = uq_teks(nm) or t.get('kunci_alami_teks') or '(tidak dinyatakan)'
    l3 = '← %s   %s   kunci alami: %s' % (asal, '[%s]' % lama if lama else '', ka)

    warna = MERAH if t.get('belum_ddl') else (AKAR_BIRU if akar else ANAK_BIRU)
    return [(l1, warna, F_KEPALA), (l2, KUNING, F_KUNCI), (l3, MUDA, F_ASAL)], warna


uq = dict(d['uq'])


def uq_teks(nm):
    u = uq.get(nm)
    return ' · '.join(' + '.join(x) for x in u) if u else ''


def kardinal(nm, fk):
    u = uq.get(nm) or []
    return any(len(x) == 1 and x[0] == fk.get('kolom') for x in u)


# ══ Lintasan pertama: siapa muncul di lembar mana (untuk penanda BERSAMA) ══
def jelajah(nm, buang, lihat, sisakan=None):
    if nm in buang or nm in lihat:
        return
    if sisakan is not None and nm not in sisakan:
        return
    lihat.add(nm)
    for a, f in anak.get(nm, []):
        jelajah(a, buang, lihat, sisakan)


for nama_l, sp in LEMBAR.items():
    lihat = set()
    for a in sp['akar']:
        if sp.get('datar'):
            lihat.add(a)
        else:
            jelajah(a, sp['buang'], lihat, sp.get('sisakan'))
    sp['isi'] = lihat
    for nm in lihat:
        di_lembar[nm].append(nama_l)

# ══ Gambar tiap lembar ═════════════════════════════════════════════════════
digambar = Counter()
for nama_l, sp in LEMBAR.items():
    ws = wb.create_sheet(nama_l)
    kepala(ws, 'TREATY IN — %s — struktur data model baru' % nama_l.upper(),
           sp['ket'], len(sp['isi']))
    r = [6]
    ditulis = set()

    def turun(nm, fk, dalam, akar=False):
        if sp.get('sisakan') is not None and nm not in sp['sisakan']:
            tolak['entitas di luar sambungan yang lembar Penyesuaian gambarkan'] += 1
            return
        if nm in sp['buang']:
            tolak['entitas dikecualikan dari lembarnya menurut daftar bernama'] += 1
            return
        if nm in ditulis:
            tolak['entitas yang sudah tergambar di lembar yang sama (tidak digambar dua kali)'] += 1
            return
        ditulis.add(nm)
        lain = [x for x in di_lembar[nm] if x != nama_l]
        b3, _ = baris_kotak(nm, fk, akar, lain)
        r[0] = kotak(ws, r[0], 2 + INDEN * dalam, b3,
                     AKAR_BIRU if akar else ANAK_BIRU)
        digambar[nama_l] += 1
        for a, f in sorted(anak.get(nm, [])):
            turun(a, f, dalam + 1)

    for a in sp['akar']:
        if sp.get('datar'):
            turun(a, None, 0, akar=True)
        else:
            turun(a, None, 0, akar=True)

# ══ Lembar Daftar Relasi ═══════════════════════════════════════════════════
ws = wb.create_sheet('Daftar Relasi')
kepala(ws, 'DAFTAR RELASI — kunci asing apa adanya',
       'sumber: ddl-usulan/ untuk 38, to-spec Adjustment untuk 3', len(d['relasi']))
JUDUL = ['#', 'Induk', 'Kard.', 'Anak', 'Kunci tamu', 'ON DELETE', 'Muncul di lembar']
LEB = [3, 26, 5, 28, 26, 20, 34]
c0 = 2
for j, (h, w) in enumerate(zip(JUDUL, LEB)):
    c = c0 + sum(LEB[:j])
    ws.merge_cells(start_row=5, start_column=c, end_row=5, end_column=c + w - 1)
    s = ws.cell(5, c, h); s.fill = AKAR_BIRU; s.font = F_TH; s.alignment = KIRI
for i, rel in enumerate(d['relasi']):
    rr = 6 + i
    lem = ', '.join(sorted(set(di_lembar[rel['anak']]) & set(di_lembar[rel['induk']]))) \
        or ', '.join(sorted(di_lembar[rel['anak']]))
    nilai = [rel.get('no', i + 1), rel['induk'], rel['kard'], rel['anak'], rel['kolom'],
             rel['on_delete'], lem]
    for j, (v, w) in enumerate(zip(nilai, LEB)):
        c = c0 + sum(LEB[:j])
        ws.merge_cells(start_row=rr, start_column=c, end_row=rr, end_column=c + w - 1)
        s = ws.cell(rr, c, v)
        s.font = F_TD
        s.alignment = KIRI
        if str(rel['on_delete']).startswith('('):
            s.fill = KUNING

# ══ Lembar Catatan & Batas ═════════════════════════════════════════════════
ws = wb.create_sheet('Catatan & Batas')
kepala(ws, 'CATATAN DAN BATAS — dibaca sebelum memakai berkas ini sebagai bukti',
       'setiap baris menyebut apa yang dapat dan tidak dapat dibuktikan', len(tabel))
CAT = [
    ('Sumber daftar entitas',
     'STRUKTUR-DATA.md — MENGIKAT, urutan wewenang butir 4. 37 di dalam gelombang ini.'),
    ('Sumber nama dan tipe kolom',
     'KAMUS-KOLOM.md — MENGIKAT, butir 5. Ia memuat 35; dua sisanya dari to-spec Adjustment.'),
    ('Sumber relasi',
     'FOREIGN KEY di ddl-usulan/ untuk 38; KEPUTUSAN-SAMBUNGAN G3 dan STRUKTUR-DATA sec 2 untuk 3.'),
    ('Sumber asal Pega',
     'peta-nama-tabel-treatyin.tsv dan SPEC-MODEL-DATA sec 2.3 / 10.22. Tertelusur 30 dari 37.'),
    ('Bentuknya ditiru dari',
     'Diagram-Skema-Tabel-TreatyIn-dan-EDM-v2.xlsx — BENTUKNYA saja. Berkas itu POTRET '
     'SISTEM LAMA 44 tabel T_; berkas ini RANCANGAN 37 entitas model baru.'),
    ('', ''),
    ('YANG DAPAT DIBUKTIKAN',
     'Bahwa sebuah entitas, kunci utama, atau kunci asing DISEBUT oleh sumber di atas.'),
    ('YANG TIDAK DAPAT DIBUKTIKAN',
     'Bahwa himpunannya LENGKAP. Semesta sec 10 masih kurang 340 properti titik buta (L-8), '
     'dan gambar dari himpunan yang bolong TERLIHAT LENGKAP.'),
    ('', ''),
    ('DUA ENTITAS BELUM BER-DDL',
     'NILAI_SELISIH dan BESARAN_DAPAT_DISESUAIKAN bertanda MERAH. Keduanya DIDAFTAR di '
     'STRUKTUR-DATA.md yang mengikat, dan NOL berkas di ddl-usulan/, NOL pasal di KAMUS-KOLOM.md. '
     'Menggambar sebuah tabel bukan membangunnya. Lihat COCOK-ENAM-SUMBER.md temuan S-1.'),
    ('ON DELETE bertanda kuning',
     'Nilai "(tidak dinyatakan)" berarti DDL-nya tidak memuatnya — temuan F-21, 38 keputusan '
     'ON DELETE yang tidak sampai ke DDL. Kuning di lembar Daftar Relasi menandainya.'),
    ('SATU PERTENTANGAN DIPILIH',
     'Berapa penunjuk versi yang disimpan NILAI_SELISIH: KEPUTUSAN-SAMBUNGAN G2 berkata DUA, '
     'STRUKTUR-DATA sec 1.1 berkata SATU. Gambar ini memakai SATU atas dasar butir 4. '
     'Itu pilihan wewenang, bukan adjudikasi — temuan S-3.'),
    ('PEMBAGIAN LEMBAR BERBEDA DARI CONTOHNYA',
     'Contohnya terbagi Prop / Non Prop / EDM Prop / EDM Non Prop. Di model baru addendum ADALAH '
     'sebuah VERSI_KONTRAK (G1), jadi tidak ada pohon EDM tersendiri. Yang tersisa: proporsional '
     'versus non-proporsional, ditambah lembar Penyesuaian dan lembar Acuan.'),
    ('Yang membuat berkas ini basi',
     'Setiap suntingan pada ddl-usulan/, STRUKTUR-DATA.md, atau KEPUTUSAN-SAMBUNGAN-ADJUSTMENT.md. '
     'Jalankan ulang rantai empat langkah di 4-erd-dan-tabel-datar/ISI-FOLDER.md.'),
]
for i, (a, b) in enumerate(CAT):
    rr = 6 + i
    ws.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=21)
    ws.merge_cells(start_row=rr, start_column=22, end_row=rr, end_column=59)
    s1 = ws.cell(rr, 2, a); s2 = ws.cell(rr, 22, b)
    s1.font = Font(name='Arial', size=8, bold=True); s1.alignment = KIRI
    s2.font = F_TD; s2.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
    if a.isupper() and a:
        s1.fill = KUNING; s2.fill = KUNING
    ws.row_dimensions[rr].height = 26 if len(b) > 110 else 14

wb.save(KELUAR)

# ══════════════════════════ CETAK ══════════════════════════════════════════
P = print
P('=== bangkitkan-struktur-data-xlsx.py ===')
P()
P('DITULIS: %s' % KELUAR)
P('  entitas di sumber : %d' % len(tabel))
for nama_l in LEMBAR:
    P('  lembar %-18s %d kotak' % (nama_l, digambar[nama_l]))
P('  lembar %-18s %d baris' % ('Daftar Relasi', len(d['relasi'])))
P('  lembar %-18s %d baris' % ('Catatan & Batas', len(CAT)))
P()
tak = [n for n in tabel if not di_lembar[n]]
P('ENTITAS yang TIDAK muncul di satu lembar pun : %d' % len(tak))
for n in tak:
    P('     - %s' % n)
P()
P('BERSAMA -- muncul di lebih dari satu lembar : %d'
  % sum(1 for n in tabel if len(di_lembar[n]) > 1))
P()
P('DITOLAK / TIDAK DIGAMBAR -- satu baris per sebab:')
for k, v in tolak.most_common():
    P('   %-66s %d' % (k, v))
if not tolak:
    P('   (tidak ada)')
P()
P('-- BATAS PERKAKAS INI --')
P('   Ia meniru BENTUK berkas contoh, bukan ISInya. Contohnya potret sistem lama')
P('   44 tabel T_; berkas ini rancangan 37 entitas model baru.')
P('   Ia TIDAK dapat menyatakan himpunannya lengkap -- L-8, 340 properti.')
sys.exit(1 if tak else 0)
